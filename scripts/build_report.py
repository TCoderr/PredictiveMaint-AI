"""Build the 3-5 page report (DOCX; PDF via scripts/export_office.ps1) from reports/metrics.json and the notebook figures.
No number is typed by hand. Run: .venv/Scripts/python scripts/build_report.py
"""
import datetime as dt
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / 'reports' / 'figures'
R = json.loads((ROOT / 'reports' / 'metrics.json').read_text(encoding='utf-8'))
AUTHOR = 'Tugay Sinan Ergün'
D, S, TD, FIN, TH, EA = R['data'], R['split'], R['test_default'], R['test_final'], R['threshold'], R['error_analysis']
SEL = R['selected_model']


def n(x, d=3): return f'{x:.{d}f}'.replace('.', ',')
def pct(x, d=1): return f'%{100 * x:.{d}f}'.replace('.', ',')
def i(x): return str(int(x))


doc = Document()
sec = doc.sections[0]; sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(1.8)
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.1
foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = foot.add_run(f'PredictiveMaint AI · {AUTHOR} · Makine ve Derin Öğrenme Uygulamaları'); fr.font.size = Pt(8)


def h(text, size=12.5):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x00, 0x79, 0x8C); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)


def para(text, style=None):
    text = text.replace('`', '')   # markdown code ticks are not rendered in Word
    p = doc.add_paragraph(style=style)
    for k, part in enumerate(text.split('**')):
        r = p.add_run(part); r.bold = k % 2 == 1
    return p


def bullets(items):
    for t in items:
        para(t, 'List Bullet').paragraph_format.space_after = Pt(1)


def table(rows, widths=None, size=8.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Light Grid Accent 1'
    for a, row in enumerate(rows):
        for b, v in enumerate(row):
            c = t.cell(a, b); c.text = str(v)
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(size)
            if widths: c.width = Cm(widths[b])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def fig(name, width, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(FIG / name), width=Cm(width))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = c.add_run(caption); r.italic = True; r.font.size = Pt(8.5)


# 1 title
t = doc.add_paragraph(); r = t.add_run('PredictiveMaint AI — Endüstriyel Ekipman Arızası Tahmini'); r.bold = True; r.font.size = Pt(17)
para(f'Makine ve Derin Öğrenme Uygulamaları · Proje raporu · {AUTHOR} (GitHub: TCoderr) · 19 Ekim 2026')
para('**Kod ve Notebook (GitHub):** https://github.com/TCoderr/PredictiveMaint-AI · '
     "**Colab'da aç:** https://colab.research.google.com/github/TCoderr/PredictiveMaint-AI/blob/main/notebooks/PredictiveMaint_AI.ipynb "
     "(notebook yerelde çalıştırıldı ve tekrarlanabilirliği doğrulandı; Colab'da çalıştırılması bu çalışmada doğrulanmadı).")

h('1. Problem tanımı')
para('Girdi olarak bir üretim sürecinin ürün tipi (L/M/H), hava ve süreç sıcaklığı, devir, tork ve takım aşınması kullanılır. '
     'Çıktı, sürecin arızayla sonuçlanıp sonuçlanmayacağıdır (**Machine failure**, ikili sınıflandırma). '
     'Arızayı önceden tahmin etmek plansız duruşları azaltır. Kaçırılan her gerçek arıza (yanlış negatif) hasar ve duruş anlamına geldiği için '
     'değerlendirmede **recall** özellikle önemlidir.')

h('2. Veri seti')
para(f"**Kaynak:** UCI Machine Learning Repository — AI4I 2020 Predictive Maintenance Dataset (S. Matzka, 2020; CC BY 4.0; doi:10.24432/C5HS5C). "
     f"{i(D['rows'])} gözlem, {i(D['cols'])} sütun; eksik değer {i(D['missing'])}, tekrarlanan satır {i(D['duplicates'])}. "
     f"Hedef: Machine failure; {i(D['failures'])} arıza ({pct(D['failure_rate'], 2)}). **Güçlü sınıf dengesizliği** var. "
     '**Veri sentetiktir:** UCI açıklamasına göre gerçek kestirimci bakım verisini yansıtmak için üretilmiştir, gerçek bir fabrikadan ölçülmemiştir.')
mw = D['mode_without_failure_by_mode']
para(f"**Arıza türü sütunları** (TWF {i(D['mode_counts']['TWF'])}, HDF {i(D['mode_counts']['HDF'])}, PWF {i(D['mode_counts']['PWF'])}, "
     f"OSF {i(D['mode_counts']['OSF'])}, RNF {i(D['mode_counts']['RNF'])}) arıza gerçekleştikten sonra bilinen bilgilerdir. "
     "Modele verilirse **hedef sızıntısı** olur; bu yüzden yalnızca hata analizinde kullanıldı. "
     f"Dosyada {i(D['failure_without_mode'])} arızanın hiçbir türü yok; {i(D['mode_without_failure'])} satırda ise tür işaretli olduğu hâlde "
     f"Machine failure = 0 (bunların {i(mw['RNF'])} tanesi RNF). "
     "UCI sayfasındaki TWF ve RNF sayıları dosyadakilerle örtüşmüyor; bu raporda dosyadaki gerçek sayılar kullanıldı. "
     '`UDI` sıra numarası ve `Product ID` ürün kimliği olduğu için çıkarıldı (Product ID\'nin ilk harfi Type ile aynıdır).')

h('3. Veri keşfi ve ön işleme')
fig('01_sinif_dagilimi.png', 14.5, 'Şekil 1 — Hedef sınıf dağılımı ve arıza türü sayıları')
rt = D['failure_rate_by_type']
bullets([
    'Arızalı süreçlerde **tork ve takım aşınması** belirgin biçimde daha yüksektir. Devir dağılımı iki uca yayılır (Şekil 2: tork–devir düzlemi). Sıcaklık farkları küçüktür.',
    f"Ürün tipine göre arıza oranı: L {n(rt['L'], 2)} %, M {n(rt['M'], 2)} %, H {n(rt['H'], 2)} %.",
    'Hava ve süreç sıcaklıkları güçlü pozitif, devir ve tork güçlü negatif ilişkilidir. Hedefle doğrusal korelasyonlar zayıftır: arıza, uç değer kombinasyonlarında yoğunlaşır.',
])
fig('03_arizayla_iliskili.png', 15.5, 'Şekil 2 — Tork–devir düzleminde arızalar ve ürün tipine göre arıza oranı')
para('**Ön işleme:**')
bullets([
    'Eksik değer yoktur.',
    'Aykırı değerler IQR kuralıyla sayıldı. Fiziksel olarak makul oldukları ve arızalar uç bölgelerde yoğunlaştığı için silinmedi.',
    'Aynı andaki ölçümlerden üç özellik türetildi: temp_diff (sıcaklık farkı), power_w (tork × açısal hız) ve strain (aşınma × tork). '
    'Bunlar gelecekteki veya arıza sonrasındaki bilgiyi kullanmaz.',
    'Type one-hot kodlandı; sayısal özellikler StandardScaler ile ölçeklendi.',
    f"Katmanlı bölme yapıldı (random_state=42): eğitim {i(S['train']['rows'])} ({i(S['train']['failures'])} arıza), "
    f"doğrulama {i(S['val']['rows'])} ({i(S['val']['failures'])}), test {i(S['test']['rows'])} ({i(S['test']['failures'])}).",
    'Bütün dönüşümler Pipeline içinde yalnızca eğitim verisine fit edildi. Bu, testlerde ölçekleyici ortalamalarıyla doğrulandı.',
])

h('4. Model seçimi ve eğitim')
rows = [['Model', 'Özellikler', 'Izgara', 'Seçilen hiperparametreler', 'Doğr. PR-AUC']]
for name, t in R['tuning'].items():
    feats = '5 sensör + Type' if 'temp_diff' not in t['features'] else '+ temp_diff, power_w, strain'
    rows.append([name, feats, i(t['grid_size']), ', '.join(f'{k}={v}' for k, v in t['best_params'].items()), n(t['val_pr_auc'], 4)])
table(rows, [3.2, 3.6, 1.2, 6.0, 2.0])
para('**Modeller:**')
bullets([
    '**Logistic Regression:** yorumlanabilir başlangıç modeli; türetilmiş özellikli bir varyantı da denendi.',
    '**Random Forest:** doğrusal olmayan eşik ilişkileri için karşılaştırma modeli; 300 ağaç.',
    '**MLPClassifier:** küçük bir sinir ağı; early stopping kullanıldı, GPU gerekmez.',
])
para('Her ızgaradaki modeller eğitim verisinde eğitildi ve **doğrulama PR-AUC** ile seçildi. PR-AUC, dengesiz veride pozitif sınıfa odaklanır. '
     f"Final model **{SEL}**, doğrulama PR-AUC değeri en yüksek olduğu için seçildi. **Test verisi hiçbir seçimde kullanılmadı.**")

h('5. Değerlendirme sonuçları (test kümesi)')
rows = [['Model', 'Accuracy', 'Precision', 'Recall', 'F1', 'ROC-AUC', 'PR-AUC', 'TP', 'FN', 'FP']]
for name, m in TD.items():
    rows.append([name + ' (0,5)', n(m['accuracy'], 4), n(m['precision']), n(m['recall']), n(m['f1']), n(m['roc_auc']), n(m['pr_auc']), i(m['tp']), i(m['fn']), i(m['fp'])])
rows.append([f'{SEL} (eşik {n(TH["value"], 2)})', n(FIN['accuracy'], 4), n(FIN['precision']), n(FIN['recall']), n(FIN['f1']), n(FIN['roc_auc']), n(FIN['pr_auc']), i(FIN['tp']), i(FIN['fn']), i(FIN['fp'])])
rows.append(["Naif 'hep arıza yok'", n(R['naive_all_zero']['accuracy'], 4), '—', '0', '—', '—', '—', '0', i(R['naive_all_zero']['fn']), '0'])
table(rows, [4.3, 1.6, 1.5, 1.4, 1.3, 1.5, 1.4, 0.9, 0.9, 0.9], size=8)
fig('05_karisiklik_matrisleri.png', 16.5, 'Şekil 3 — Test kümesi karışıklık matrisleri (eşik 0,5)')
para(f"**Accuracy neden yeterli değil?** Her zaman \"arıza yok\" diyen naif model, hiçbir arızayı yakalamadığı hâlde {n(R['naive_all_zero']['accuracy'], 4)} accuracy alır. "
     f"Başlangıç modeli LogReg (temel) {n(TD['LogReg (temel)']['accuracy'], 4)} accuracy'ye rağmen arızaların yalnızca {n(TD['LogReg (temel)']['recall'])} oranını yakalar. "
     'Dengesiz problemde başarıyı precision, recall, F1 ve PR-AUC gösterir.')

h('6. Yanlış negatifler ve karar eşiği')
v5, vt, d5 = TH['val_at_0_5'], TH['val_at_threshold'], TD[SEL]
para(f"**Eşik seçimi:** {SEL} için karar eşiği **yalnızca doğrulama verisinde** F2 skorunu en yükseğe çıkaran değer olarak seçildi: **{n(TH['value'], 2)}**. "
     f"F2, recall'a iki kat ağırlık verir. Doğrulamada eşik 0,5 iken recall {n(v5['recall'])} ve precision {n(v5['precision'])}; "
     f"seçilen eşikte recall {n(vt['recall'])} ve precision {n(vt['precision'])}. Bu eşik test kümesinde **bir kez** uygulandı.")
bullets([
    f"Testte {i(FIN['tp'] + FIN['fn'])} gerçek arızadan **{i(FIN['fn'])} tanesi kaçırıldı (FN)**. Eşik 0,5'te bu sayı {i(d5['fn'])} idi.",
    f"Yanlış negatif oranı FNR = FN / (FN + TP) = **{n(FIN['fnr'])}** = 1 − recall ({n(FIN['recall'])}). Recall arttıkça kaçırılan arıza oranı aynı ölçüde azalır.",
    f"Bedeli: yanlış alarm (FP) sayısı {i(d5['fp'])}'den {i(FIN['fp'])}'e çıktı, precision {n(d5['precision'])}'ten {n(FIN['precision'])}'e düştü.",
    'Bakımda bir arızayı kaçırmak ekipman hasarı, plansız duruş ve güvenlik riski yaratabilir. Yanlış alarm ise genellikle gereksiz bir kontrol demektir. '
    'Doğru denge gerçek maliyetlere bağlıdır. Gerçek işletme maliyeti bilinmediği için **parasal maliyet hesabı yapılmadı**.',
])
fig('07_esik_dengesi.png', 11.5, 'Şekil 4 — Doğrulama kümesinde eşiğe göre precision, recall ve F2')

h('7. Açıklanabilirlik')
coef = R['explain']['logreg_eng_coef']; perm = R['explain']['rf_permutation_pr_auc_drop']
top_perm = sorted(perm, key=perm.get, reverse=True)[:4]
bullets([
    f"**LogReg katsayıları (standartlaştırılmış):** En büyük pozitif katsayılar Torque [Nm] ({n(coef['Torque [Nm]'], 2)}) ve Rotational speed "
    f"({n(coef['Rotational speed [rpm]'], 2)}), en büyük negatif katsayı power_w ({n(coef['power_w'], 2)}). "
    'Bu üç özellik birbirine bağlı olduğu için (güç = tork × devir) katsayılar tek başına etki büyüklüğü olarak okunmamalıdır. '
    'Type one-hot katsayıları birbirine göre yorumlanır.',
    f"**Random Forest permutation importance** (doğrulama PR-AUC düşüşü): en önemli özellikler {', '.join(top_perm)}. "
    'Bu, modelin neye dayandığını gösterir; **nedensellik kanıtı değildir**. İlişkili özellikler önemi aralarında paylaşabilir.',
])
fig('09_ozellik_onemleri.png', 15.5, 'Şekil 5 — LogReg katsayıları ve Random Forest permutation importance')

h('8. Hata analizi')
fbm, fnm = EA['failures_by_mode'], EA['fn_by_mode']
rows = [['Arıza türü (test)', 'Arıza', 'Kaçırılan (FN)']] + [[k, i(fbm[k]), i(fnm[k])] for k in fbm]
table(rows, [5, 2.5, 3])
fpm, tnm = EA['fp_vs_tn_means']['FP ortalaması'], EA['fp_vs_tn_means']['TN ortalaması']
hard = [k for k, v in fnm.items() if v > 0]
bullets([
    f"Kaçırılan arızaların tamamı şu gruplarda: **{', '.join(hard)}**. HDF, PWF ve OSF arızalarının hepsi yakalandı.",
    'UCI açıklamasına göre TWF, takım aşınması 200–240 dakika arasındayken rastgele bir anda gerçekleşir. "Tür yok" arızalar ise etiket tutarsızlığıdır. '
    'Bu nedenle bu gruplar mevcut özelliklerle ayırt edilemez. '
    f"Test kümesinde RNF etiketli arıza yoktur.",
    f"Yanlış alarm sayısı {i(EA['fp_count'])}; hiçbirinde arıza türü işareti yok ({i(EA['fp_with_failure_mode_flag'])}). "
    f"Yanlış alarmlarda ortalama tork {n(fpm['Torque [Nm]'], 1)} (doğru negatiflerde {n(tnm['Torque [Nm]'], 1)}), aşınma {n(fpm['Tool wear [min]'], 1)} "
    f"({n(tnm['Tool wear [min]'], 1)}): model, arızalı süreçlere benzeyen uç bölgelerde alarm veriyor.",
])

h('9. Sonuç ve geliştirme önerileri')
para(f"Türetilmiş özelliklerle eğitilen {SEL}, başlangıç modeline göre çok daha iyi sonuç verdi. "
     f"Test PR-AUC: {n(TD[SEL]['pr_auc'])}'e karşı {n(TD['LogReg (temel)']['pr_auc'])}. "
     f"Recall eşik 0,5'te {n(d5['recall'])}, doğrulamayla seçilen eşikte {n(FIN['recall'])} oldu. "
     '**Sınırlar:**')
bullets([
    'Veri sentetiktir ve arıza kuralları bilinen formüllerle üretilmiştir. Türetilmiş özellikler bu kuralları yansıttığı için gerçek veride performans belirgin biçimde düşebilir.',
    'Zaman boyutu yoktur; model arızadan önce uyarı üretmez, sürecin arızalı olup olmadığını sınıflandırır.',
    'Model ve eşik tek bir bölmeyle seçilmiştir.',
    'Bu çalışma **akademik bir prototiptir**; gerçek fabrikalarda doğrudan güvenilir şekilde kullanılabileceği iddia edilmez.',
])
para('**Öneriler:**')
bullets([
    'Gerçek ve zaman damgalı sensör verisiyle yeniden eğitmek.',
    'Tabakalı çapraz doğrulamayla model ve eşik seçmek.',
    'Maliyet bilgisi varsa eşiği maliyete göre belirlemek.',
    'Olasılık kalibrasyonu yapmak.',
    'Arıza türüne göre çok sınıflı modeller denemek.',
])
para(f"**Ortam:** Python {R['versions']['python']}, scikit-learn {R['versions']['sklearn']}, pandas {R['versions']['pandas']}. "
     "Geliştirme Claude Code (Anthropic) yardımıyla yapıldı; proje sahibi ve sorumlusu Tugay Sinan Ergün'dür.")

cp = doc.core_properties; cp.author = cp.last_modified_by = AUTHOR; cp.title = 'PredictiveMaint AI Proje Raporu'
cp.created = cp.modified = dt.datetime(2026, 10, 19, 12, 0)
out = ROOT / 'reports' / 'PredictiveMaint_AI_Rapor.docx'; doc.save(out); print(out)
